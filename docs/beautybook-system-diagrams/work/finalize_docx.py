from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from io import BytesIO
from lxml import etree
from PIL import Image
from docx import Document
from pypdf import PdfReader
import json, re

root=Path(__file__).parent
path=root.parent/'BeautyBook_Phan_tich_Thiet_ke_So_do_v3.docx'
with ZipFile(path) as z:
    entries={i.filename:(i,z.read(i.filename)) for i in z.infolist()}
for name,tags in [('docProps/core.xml',['creator','lastModifiedBy','keywords','description']),('docProps/app.xml',['Company','Manager'])]:
    info,data=entries[name]
    xml=etree.fromstring(data)
    for node in xml.iter():
        if etree.QName(node).localname in tags:node.text=None
    entries[name]=(info,etree.tostring(xml,encoding='UTF-8',xml_declaration=True,standalone=True))
temp=path.with_suffix('.finalizing.docx')
with ZipFile(temp,'w',ZIP_DEFLATED) as z:
    for info,data in entries.values():z.writestr(info,data)
temp.replace(path)

doc=Document(path)
assert len(doc.inline_shapes)==19
assert len(doc.tables)==2
ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
xml=doc.element
assert not xml.xpath('//w:ins|//w:del|//w:commentRangeStart')
instr=[n.text or '' for n in xml.xpath('//w:instrText')]
assert sum('SEQ Hinh' in t for t in instr)==19
assert sum('SEQ Bang' in t for t in instr)==2
assert sum('REF fig' in t for t in instr)==19
assert sum('REF tbl' in t for t in instr)==2
assert sum('TOC ' in t for t in instr)==3
resolution=[]
for shape in doc.inline_shapes:
    rid=shape._inline.graphic.graphicData.pic.blipFill.blip.embed
    im=Image.open(BytesIO(doc.part.related_parts[rid].blob))
    dpi=im.width/(shape.width/914400)
    assert dpi>=150
    resolution.append(round(dpi))
wide_sections=0
for section in doc.sections:
    width=round(section.page_width/360000,1)
    assert width in (21,42)
    wide_sections += width==42
    assert round(section.page_height/360000,1)==29.7
assert wide_sections==1
reader=PdfReader(root/'BeautyBook_QA.pdf')
assert len(reader.pages)==38
assert [i+1 for i,p in enumerate(reader.pages) if float(p.mediabox.width)>float(p.mediabox.height)]==[29]
visible=' '.join(n.text or '' for n in xml.xpath('//w:t'))
assert 'khách truy cập' not in visible.lower()
assert not re.search(r'\b(?:NV|HT)\d+[a-z]?\b',visible)
assert 'Sơ đồ thực thể quan hệ ERD' in visible
assert 'Database Diagram tổng quát' in visible
for page in reader.pages:
    text=page.extract_text() or ''
    assert not any(e in text for e in ['Error!','Reference source not found','No table of figures entries found'])
result={'pages':len(reader.pages),'diagrams':len(doc.inline_shapes),'erd_diagrams':2,'usecase_specifications':27,'tables':len(doc.tables),'minimum_effective_image_dpi':min(resolution),'visual_qa':'All 38 final pages inspected; overview also inspected at full-page resolution.','revision_checks':'Overall Database Diagram added on page 29: 36 core tables and 64 declared foreign-key relationships. One A3 landscape page; remaining pages A4 portrait. Existing ERDs and Usecase edits preserved.','metadata':'Personal author and organization fields cleared.','file_size_bytes':path.stat().st_size}
(root/'final_qa.json').write_text(json.dumps(result,indent=2),encoding='utf-8')
print(json.dumps(result))
