from doc_helpers import Report, ROOT, Cm, Pt, WD_ALIGN_PARAGRAPH, WD_SECTION_START
from specifications import BUSINESS, SYSTEM, SUBCASES
from pathlib import Path
from docx.enum.section import WD_ORIENT

r=Report();d=r.doc
p=d.add_paragraph('BEAUTYBOOK','Title');p.paragraph_format.space_before=Pt(110);p.paragraph_format.space_after=Pt(20)
p=d.add_paragraph('PHÂN TÍCH VÀ THIẾT KẾ HỆ THỐNG','Title');p.paragraph_format.space_after=Pt(20)
d.add_paragraph('Hệ thống đặt lịch dịch vụ làm đẹp','Subtitle')
p=r.p('Usecase nghiệp vụ và hệ thống\nĐặc tả Usecase\nSơ đồ lớp phân tích và thiết kế\nSơ đồ ERD và mô hình dữ liệu');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.space_before=Pt(26);p.paragraph_format.line_spacing=1.5
p=r.p('Tài liệu phục vụ báo cáo đồ án');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.space_before=Pt(45)
p=r.p('Phiên bản ngày 15 tháng 10 năm 2026');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Cm(0)
d.add_section(WD_SECTION_START.NEW_PAGE).different_first_page_header_footer=False
p=d.add_paragraph('MỤC LỤC','Title');r.field(d.add_paragraph(),'TOC \\o "1-3" \\h \\z \\u','')
r.page();d.add_paragraph('DANH MỤC HÌNH VÀ BẢNG','Title')
r.label('Danh mục hình').runs[0].bold=True;r.field(d.add_paragraph(),'TOC \\h \\z \\c "Hinh"','')
r.label('Danh mục bảng').runs[0].bold=True;r.field(d.add_paragraph(),'TOC \\h \\z \\c "Bang"','')

r.heading('Phân tích hệ thống',1,True)
r.heading('Phạm vi và tác nhân',2)
r.p('Trong đồ án BeautyBook, em xây dựng mô hình hệ thống theo hai góc nhìn: hoạt động cung cấp dịch vụ làm đẹp và các chức năng mà phần mềm hỗ trợ. Phần phân tích xác định mục tiêu của từng tác nhân; phần thiết kế trình bày các lớp xử lý và cấu trúc dữ liệu phục vụ những mục tiêu đó.')
r.p('Phạm vi tài liệu tập trung vào website quản lý cơ sở và quản trị hệ thống, cùng ứng dụng dành cho khách hàng. Các mô hình bao quát hồ sơ cơ sở, dịch vụ, nhân sự, đặt lịch, định kỳ, ưu đãi, đánh giá, thông báo và thống kê. Những bảng kỹ thuật như phiên đăng nhập, hàng đợi thông báo và nhật ký chi tiết được lược khỏi hình dữ liệu cốt lõi để giữ sơ đồ dễ đọc.')
r.table(['Tác nhân','Vai trò trong BeautyBook'],[
 ['Khách hàng','Đăng ký, đăng nhập, tìm kiếm và xem thông tin cơ sở; đặt lịch, theo dõi lịch, gửi yêu cầu thay đổi và đánh giá trải nghiệm.'],
 ['Chủ cơ sở','Quản lý doanh nghiệp, chi nhánh, dịch vụ, ưu đãi, nhân sự và hoạt động thuộc quyền sở hữu.'],
 ['Lễ tân','Tiếp nhận và điều phối lịch, xác nhận khách đến trong chi nhánh được cấp quyền.'],
 ['Chuyên viên','Xem lịch và dịch vụ được phân công; không cập nhật tiến trình lịch trong luồng thông thường.'],
 ['Quản trị hệ thống','Xét duyệt cơ sở, quản trị tài khoản và danh mục, kiểm duyệt nội dung, giám sát toàn nền tảng.']],[4,12],'Các tác nhân tham gia hệ thống')
r.p('Trong sơ đồ nghiệp vụ, “Cơ sở làm đẹp” đại diện cho tổ chức cung cấp dịch vụ; khi chuyển sang sơ đồ hệ thống, tác nhân này được tách thành chủ cơ sở, lễ tân và chuyên viên. Quyền truy cập còn phụ thuộc vào doanh nghiệp và chi nhánh được cấp, không chỉ dựa trên tên vai trò.')

r.heading('Sơ đồ Usecase nghiệp vụ',2,True)
r.figure('01_usecase_nghiep_vu','Usecase nghiệp vụ BeautyBook','Các Usecase nghiệp vụ mô tả kết quả mà khách hàng, cơ sở và nền tảng cần đạt được.',
 'Các đường nối biểu diễn sự tham gia vào nghiệp vụ, không biểu diễn trình tự thực hiện. Công việc phục vụ trực tiếp tại cơ sở có thể diễn ra ngoài phần mềm.')
r.heading('Đặc tả Usecase nghiệp vụ',2,True)
for uc in BUSINESS:r.spec(uc)

r.heading('Sơ đồ Usecase hệ thống',2,True)
r.figure('02_usecase_khach_hang','Usecase hệ thống dành cho khách hàng','Các chức năng trên kênh khách hàng được nhóm theo mục tiêu sử dụng.',
 'Khách hàng có thể tra cứu thông tin công khai trước khi đăng nhập. Các chức năng đặt lịch và quản lý dữ liệu cá nhân yêu cầu tài khoản phù hợp. Đăng nhập được xem là tiền điều kiện, không nối include từ mọi Usecase.')
r.heading('Chức năng vận hành tại cơ sở',3,True)
r.figure('03_usecase_co_so','Usecase hệ thống dành cho cơ sở','Nhóm chức năng vận hành phân biệt quyền của chủ cơ sở, lễ tân và chuyên viên.',
 'Chuyên viên chỉ xem lịch được phân công. Chủ cơ sở và lễ tân thực hiện check-in trong Usecase Tiếp nhận và xử lý lịch hẹn. Các chức năng đăng nhập dùng chung được lược khỏi hình này.')
r.heading('Chức năng quản trị nền tảng',3,True)
r.figure('04_usecase_admin','Usecase hệ thống dành cho quản trị','Quản trị hệ thống thực hiện xét duyệt, quản trị dữ liệu và giám sát trong phạm vi quyền được cấp.',
 'Quản lý danh mục và nội dung bao gồm danh mục dịch vụ chuẩn và nội dung cần kiểm duyệt. Thống kê và tra cứu toàn hệ thống hỗ trợ xem cơ sở theo nhóm dịch vụ, người dùng và lịch hẹn theo các tiêu chí phù hợp.')
r.heading('Phân rã chức năng đặt lịch',3,True)
r.figure('05_usecase_dat_lich','Phân rã Usecase đặt lịch hẹn','Đặt lịch hẹn luôn cần kiểm tra giờ trống, tính giá và ghi nhận lịch.',
 'Mũi tên nét đứt «include» hướng từ Đặt lịch hẹn đến hành vi bắt buộc được sử dụng. Áp dụng voucher là lựa chọn bên trong bước tính giá; đánh giá là mục tiêu độc lập với đặt lịch, vì vậy không nối extend giữa hai chức năng này.')

r.heading('Đặc tả Usecase hệ thống',2,True)
r.p('Mỗi đặc tả nêu mục tiêu, tác nhân, điều kiện trước, luồng chính, ngoại lệ và điều kiện sau. Các Usecase được gọi trực tiếp bằng tên chức năng. Ba đặc tả Kiểm tra giờ trống, Tính giá và ưu đãi, Ghi nhận lịch hẹn trình bày các bước bắt buộc trong đặt lịch.')
for uc in SYSTEM:r.spec(uc)
for uc in SUBCASES:r.spec(uc)

r.heading('Sơ đồ lớp phân tích',2,True)
r.p('Ở mức phân tích, em sử dụng tên lớp theo nghiệp vụ và chỉ giữ các đặc điểm cần để hiểu đối tượng. Hai hình dưới đây cùng mô tả một mô hình; lớp Nhân viên xuất hiện lại để nối phần danh mục với phần lịch hẹn. Bội số đặt ở đầu nào cho biết số đối tượng có thể xuất hiện ở đầu đó.')
r.figure('06_lop_phan_tich_danh_muc','Lớp phân tích về cơ sở và danh mục','Doanh nghiệp sở hữu danh mục dịch vụ, còn chi nhánh là nơi trực tiếp cung cấp dịch vụ.',
 'Quan hệ nhiều–nhiều giữa Nhân viên và Dịch vụ tại chi nhánh biểu diễn năng lực được gán. Khi thiết kế dữ liệu, quan hệ này được tách bằng bảng liên kết staff_services.')
r.heading('Các lớp liên quan đến lịch hẹn',3,True)
r.figure('07_lop_phan_tich_lich_hen','Lớp phân tích về lịch hẹn và đánh giá','Lịch hẹn kết nối nhu cầu của khách với các dịch vụ và chuyên viên phục vụ.',
 'Một lịch được tạo hợp lệ phải có ít nhất một dòng dịch vụ. Lịch đơn không bắt buộc thuộc kế hoạch định kỳ. Một dòng dịch vụ có thể chưa được gán chuyên viên, còn mỗi lịch có tối đa một đánh giá.')

r.heading('Thiết kế hệ thống',1,True)
r.heading('Sơ đồ lớp thiết kế',2)
r.p('Phần thiết kế chuyển các nhiệm vụ nghiệp vụ thành lớp tiếp nhận yêu cầu, lớp xử lý và lớp truy cập dữ liệu. Sơ đồ dưới đây chọn các lớp cốt lõi và một số phương thức đại diện. Ký hiệu “+” biểu diễn phương thức công khai; mũi tên nét đứt biểu diễn quan hệ sử dụng.')
r.figure('08_lop_thiet_ke_danh_muc','Lớp thiết kế quản lý cơ sở và danh mục','Các Controller chuyển yêu cầu hợp lệ đến Service phụ trách.',
 'Các Service truy vấn dữ liệu qua PrismaService. StaffService là lớp xử lý nghiệp vụ; staff_services là bảng liên kết nhân viên với dịch vụ.',width=14.8)
r.heading('Các lớp xử lý đặt lịch',3,True)
r.figure('09_lop_thiet_ke_dat_lich','Lớp thiết kế đặt lịch định kỳ và đánh giá','RecurringService sử dụng BookingsService để tạo từng kỳ; BookingsService sử dụng PricingEngineService để xử lý giá.',
 'PrismaService kế thừa PrismaClient. Hình chỉ vẽ một số đường phụ thuộc chính; các lớp còn phối hợp với phân quyền, thiết lập nền tảng, thông báo và xử lý thay đổi lịch. Chuỗi lịch được tạo từng kỳ và có cơ chế hoàn tác khi lỗi, không được coi là một giao dịch cơ sở dữ liệu duy nhất.')

r.heading('Sơ đồ thực thể quan hệ ERD',2,True)
r.p('Em chia ERD cốt lõi thành hai phần liên kết qua Dịch vụ tại chi nhánh và Nhân viên. Hình chữ nhật biểu diễn thực thể, hình thoi biểu diễn quan hệ; bội số ghi ở mỗi đầu cho biết số đối tượng có thể có tại đầu đó. Thuộc tính, khóa chính và khóa ngoại được trình bày trong phần mô hình dữ liệu.')
r.figure('17_erd_co_so_dich_vu','ERD cơ sở và dịch vụ','Phần thứ nhất mô tả cách tổ chức cơ sở, danh mục dịch vụ và năng lực nhân viên.',
 'Một doanh nghiệp có nhiều chi nhánh và dịch vụ trong danh mục. Mỗi dịch vụ tại chi nhánh thuộc một chi nhánh và một dịch vụ doanh nghiệp. Quan hệ nhiều–nhiều giữa nhân viên và dịch vụ được triển khai bằng bảng staff_services.',width=14.5)
r.heading('ERD đặt lịch và đánh giá',3,True)
r.figure('18_erd_dat_lich_danh_gia','ERD đặt lịch và đánh giá','Phần thứ hai mô tả quan hệ giữa khách hàng, lịch hẹn, dịch vụ, người thực hiện và đánh giá.',
 'Mỗi lịch hợp lệ có ít nhất một chi tiết dịch vụ và tối đa một đánh giá. Mỗi chi tiết gắn với một dịch vụ tại chi nhánh, có thể chưa được gán nhân viên. Ký hiệu 1 là đúng một; 0..1 là tối đa một; 0..* là không hoặc nhiều; 1..* là ít nhất một.')

sec=d.add_section(WD_SECTION_START.NEW_PAGE)
sec.orientation=WD_ORIENT.LANDSCAPE
sec.page_width=Cm(42);sec.page_height=Cm(29.7)
sec.different_first_page_header_footer=False
r.heading('Database Diagram tổng quát',2)
r.figure('19_database_diagram_tong_quat','Database Diagram tổng quát các bảng cốt lõi',
 'Em đặt nhóm lịch hẹn ở trung tâm để liên kết khách hàng, cơ sở, dịch vụ và kết quả phục vụ. Sơ đồ chọn 36 bảng và 64 FK chính; các FK bổ sung xem ở phần chi tiết. Tên bảng theo cơ sở dữ liệu, tên thuộc tính theo mô hình Prisma.',width=37)
sec=d.add_section(WD_SECTION_START.NEW_PAGE)
sec.orientation=WD_ORIENT.PORTRAIT
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.different_first_page_header_footer=False

r.heading('Mô hình dữ liệu',2)
r.p('Bảy hình sau sử dụng tên bảng vật lý và tên thuộc tính theo mô hình Prisma. PK là khóa chính; FK là khóa ngoại; UQ là ràng buộc duy nhất; “?” là trường có thể để trống. Bội số ghi trên đường nối; các FK sang hình khác vẫn được ghi trong hộp.')
dbfigs=[
 ('10_du_lieu_tai_khoan','Dữ liệu tài khoản và doanh nghiệp','Tài khoản có thể gắn với hồ sơ khách hàng hoặc chủ cơ sở.','Một tài khoản có tối đa một hồ sơ khách hàng và một hồ sơ chủ cơ sở. user_roles lưu vai trò theo phạm vi doanh nghiệp hoặc chi nhánh; đây là liên kết cấp quyền, không phải quan hệ kế thừa giữa các loại tài khoản.'),
 ('11_du_lieu_danh_muc','Dữ liệu chi nhánh và danh mục dịch vụ','Dịch vụ được tổ chức theo danh mục doanh nghiệp và dịch vụ thực tế ở chi nhánh.','canonical_services là nhóm dịch vụ chuẩn của nền tảng; business_services là danh mục doanh nghiệp; services là dịch vụ được cung cấp tại chi nhánh. Cặp branchId và businessServiceId trong services là duy nhất. service_variants lưu các lựa chọn của dịch vụ tại chi nhánh.'),
 ('12_du_lieu_nhan_su','Dữ liệu nhân viên và khả năng phục vụ','Nhân viên có hồ sơ, năng lực dịch vụ và phạm vi chi nhánh.','staff_profiles.userId có thể để trống và là duy nhất khi có giá trị. staff_branch_assignments ghi phân công chi nhánh theo thời gian. staff_services có khóa chính ghép staffId và serviceId; bảng này không thay thế việc phân công người thực hiện cho từng lịch hẹn.'),
 ('13_du_lieu_lich_hen','Dữ liệu lịch hẹn và chuỗi định kỳ','bookings là dữ liệu lịch hẹn; booking_services lưu từng dòng dịch vụ trong lịch.','Các FK customerId và branchId của bookings trỏ đến customer_profiles và branches. Mỗi lịch có tối đa một booking_contacts và có thể có nhiều payments. Mỗi lịch được tạo hợp lệ có ít nhất một booking_services, dù quan hệ cơ sở dữ liệu cho phép 0..*. serviceIds trong kế hoạch định kỳ là JSON; variantId trong dòng lịch là tham chiếu logic, không phải FK được khai báo.'),
 ('14_du_lieu_combo_khuyen_mai','Dữ liệu combo và khuyến mãi','Combo liên kết nhiều dịch vụ, còn khuyến mãi có phạm vi áp dụng riêng.','combo_services liên kết combos với services và lưu số lượng, thứ tự. Các bảng promotion_branches, promotion_services và promotion_combos có khóa chính ghép. businessId của promotions thể hiện doanh nghiệp sở hữu; phạm vi áp dụng theo doanh nghiệp được lưu riêng ở promotion_businesses.'),
 ('15_du_lieu_voucher','Dữ liệu voucher và kết quả tính giá','Voucher cấp cho khách và lượt sử dụng được lưu thành các bản ghi riêng.','customer_vouchers liên kết khách với voucher; voucher_redemptions ghi nhận lượt sử dụng theo lịch. pricing_snapshots lưu kết quả tính giá tại lúc đặt và có tối đa một bản ghi cho mỗi lịch. usedBookingId trong customer_vouchers là tham chiếu logic. promotion_businesses nối promotions với businesses và được đặt cùng hình để hoàn thiện nhóm phạm vi ưu đãi.'),
 ('16_du_lieu_danh_gia_thong_bao','Dữ liệu đánh giá và thông báo','Đánh giá chi tiết gắn với dòng dịch vụ thực tế trong lịch.','Một booking có tối đa một review; một booking_service có tối đa một review_service_rating. staffId của điểm đánh giá có thể để trống. notifications thuộc về một user và có thể liên kết với một booking qua relatedBookingId.')]
for i,(name,title,intro,note) in enumerate(dbfigs):
    r.heading(title,3,i>0)
    r.figure(name,title,intro,note,width=14.5 if i==0 else 16)

r.heading('Quy tắc liên kết và toàn vẹn dữ liệu',2,True)
r.table(['Nhóm','Quy tắc cần bảo đảm'],[
 ['Cơ sở và quyền','Chi nhánh thuộc một doanh nghiệp. Mọi thao tác quản lý phải kiểm tra đúng doanh nghiệp và chi nhánh được cấp quyền.'],
 ['Dịch vụ','serviceId trong lịch, combo, phân công nhân viên và ưu đãi đều trỏ đến services. Một doanh nghiệp có thể cung cấp cùng dịch vụ tại nhiều chi nhánh.'],
 ['Lịch hẹn','customerId và branchId là bắt buộc. Chi tiết lịch lưu giá, thời lượng và tên tại lúc đặt, giúp lịch sử không bị đổi theo giá danh mục mới.'],
 ['Chuyên viên','staff_services xác định khả năng nhận dịch vụ; booking_services.staffId xác định người được giao trong một lịch cụ thể.'],
 ['Định kỳ','Mỗi kỳ là một booking riêng. Lỗi khi tạo chuỗi phải xử lý các kỳ đã tạo và trả kết quả rõ ràng.'],
 ['Đánh giá','Đánh giá và điểm chi tiết phải thuộc đúng khách, lịch, dòng dịch vụ và chuyên viên liên quan.'],
 ['Ưu đãi','Phạm vi áp dụng, thời hạn, số lượt và điều kiện khách phải được kiểm tra lại lúc đặt; thấy ưu đãi không bảo đảm được giảm giá.']],[3.2,12.8],'Các quy tắc liên kết dữ liệu cốt lõi')
r.p('Các FK riêng lẻ không tự bảo đảm tính nhất quán giữa mọi trường. Ví dụ dịch vụ, danh mục doanh nghiệp và chuyên viên của một dòng lịch còn phải được kiểm tra để bảo đảm cùng phạm vi phục vụ. Một số ràng buộc thuộc tầng nghiệp vụ và phải được kiểm tra trước khi ghi dữ liệu.')
r.p('Yêu cầu thay đổi lịch được lưu trong appointment_change_requests, liên kết với bookings, người yêu cầu và người xử lý; chuyên viên được đề nghị có thể để trống. Sơ đồ tổng quát thể hiện quan hệ với lịch hẹn và chuyên viên được đề nghị. Các bảng ảnh liên kết đối tượng với media_files. Nhóm audit, ảnh, phân quyền chi tiết và phạm vi voucher được lược khỏi các hình dữ liệu để tập trung vào cấu trúc cốt lõi; việc lược hình không thay đổi vai trò của chúng trong hệ thống.')

r.heading('Các quy ước nghiệp vụ cần thống nhất',2,True)
for title,text in [
 ('Tiếp nhận và lịch sử','Luồng thông thường ghi nhận khách đến bằng CHECKED_IN. Khi đã qua giờ kết thúc dự kiến, lịch được phân loại vào lịch sử ở dữ liệu hiển thị. Điều này không tự cập nhật trạng thái COMPLETED, không tạo thời gian hoàn thành thực tế và không chứng minh đã thu tiền.'),
 ('Điều kiện đánh giá','Khách chỉ đánh giá lịch của mình khi đủ điều kiện và chưa gửi đánh giá trước đó. Điều kiện có thể đáp ứng với lịch COMPLETED hoặc CHECKED_IN đã qua giờ dự kiến; điểm từng dòng dịch vụ được nhập độc lập với điểm tổng thể.'),
 ('Hủy và thay đổi lịch','Hủy trực tiếp, yêu cầu hủy sát giờ, đổi giờ và đổi chuyên viên phải được xử lý theo trạng thái cùng chính sách hiện hành. Không coi mọi lịch đã qua giờ hoặc mọi yêu cầu chưa xử lý là lịch hoàn thành.'),
 ('Vai trò và phạm vi','Chuyên viên chỉ xem lịch được phân công. Chủ cơ sở và lễ tân vận hành trong phạm vi quyền của mình. Tài khoản khách hàng và dữ liệu tiếp nhận khách tại quầy không được tự đồng nhất chỉ dựa trên số điện thoại.'),
 ('Giới hạn của các mô hình','Các hình mô tả cấu trúc và hành vi cốt lõi của BeautyBook tại thời điểm lập tài liệu. Cổng thanh toán trực tuyến không được đưa thành tác nhân trong bộ Usecase; phần dữ liệu payments thể hiện khả năng lưu giao dịch, không khẳng định đã tích hợp cổng thanh toán hoặc đã triển khai đầy đủ giao diện tài chính.')]:
    r.heading(title,3);r.p(text)

output=ROOT.parent/'BeautyBook_Phan_tich_Thiet_ke_So_do_v3.docx'
r.save(output)
print(output)
print('Figures:',len(r.figures),'Tables:',len(r.tables),'Usecase specs:',len(BUSINESS)+len(SYSTEM)+len(SUBCASES))
