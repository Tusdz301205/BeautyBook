from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).parent
class Report:
    def __init__(self):
        self.doc=Document(); self.figures=[]; self.tables=[]; self.bm=100
        sec=self.doc.sections[0]
        sec.page_width=Cm(21);sec.page_height=Cm(29.7)
        sec.top_margin=Cm(2);sec.bottom_margin=Cm(2);sec.left_margin=Cm(3);sec.right_margin=Cm(2)
        sec.header_distance=Cm(.8);sec.footer_distance=Cm(.8)
        sec.different_first_page_header_footer=True
        for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption','List Number','List Bullet','Header','Footer']:
            s=self.doc.styles[name];s.font.name='Times New Roman';s.font.color.rgb=RGBColor(0,0,0)
            rp=s.element.get_or_add_rPr()
            fonts=rp.find(qn('w:rFonts'))
            for attr in list(fonts.attrib):
                if 'theme' in attr.lower():del fonts.attrib[attr]
            for slot in ['ascii','hAnsi','eastAsia','cs']:fonts.set(qn('w:'+slot),'Times New Roman')
            color=rp.find(qn('w:color'))
            for attr in list(color.attrib):
                if 'theme' in attr.lower():del color.attrib[attr]
            pp=s.element.get_or_add_pPr()
            for border in list(pp.findall(qn('w:pBdr'))):pp.remove(border)
            s.element.get_or_add_rPr().append(self.el('w:lang',{'w:val':'vi-VN'}))
            s.paragraph_format.widow_control=True
        n=self.doc.styles['Normal'];n.font.size=Pt(13)
        n.paragraph_format.line_spacing=1.15;n.paragraph_format.space_after=Pt(5);n.paragraph_format.first_line_indent=Cm(1)
        n.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
        for name in ['Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption','List Number','List Bullet','Header','Footer']:
            self.doc.styles[name].paragraph_format.first_line_indent=Cm(0)
        for i in range(1,4):
            s=self.doc.styles[f'Heading {i}'];s.font.size=Pt(14 if i==1 else 13);s.font.bold=True;s.font.italic=i==3
            s.paragraph_format.space_before=Pt(10);s.paragraph_format.space_after=Pt(6);s.paragraph_format.keep_with_next=True
            s.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
        title=self.doc.styles['Title'];title.font.size=Pt(18);title.font.bold=True;title.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
        self.doc.styles['Subtitle'].font.size=Pt(14);self.doc.styles['Subtitle'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
        cap=self.doc.styles['Caption'];cap.font.size=Pt(12);cap.font.italic=True;cap.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER;cap.paragraph_format.space_after=Pt(8)
        for name in ['Header','Footer']:
            self.doc.styles[name].font.size=Pt(10 if name=='Header' else 11)
            self.doc.styles[name].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
        self.heading_numbering()
        p=sec.header.paragraphs[0];p.style='Header';p.text='BeautyBook   |   Phân tích và thiết kế hệ thống'
        borders=self.el('w:pBdr');borders.append(self.el('w:bottom',{'w:val':'single','w:sz':'4','w:color':'808080'}));p._p.get_or_add_pPr().append(borders)
        p=sec.footer.paragraphs[0];p.style='Footer';self.field(p,'PAGE','1')
        settings=self.doc.settings.element
        settings.append(self.el('w:updateFields',{'w:val':'true'}))
        self.doc.core_properties.title='Phân tích và thiết kế hệ thống BeautyBook'
        self.doc.core_properties.subject='Usecase và đặc tả sơ đồ lớp mô hình dữ liệu'
        self.doc.core_properties.author='';self.doc.core_properties.last_modified_by='';self.doc.core_properties.comments=''
    @staticmethod
    def el(tag,attrs={}):
        node=OxmlElement(tag)
        for k,v in attrs.items():node.set(qn(k),v)
        return node
    def heading_numbering(self):
        root=self.doc.part.numbering_part.element
        abstract=self.el('w:abstractNum',{'w:abstractNumId':'70'});abstract.append(self.el('w:multiLevelType',{'w:val':'multilevel'}))
        for i in range(3):
            level=self.el('w:lvl',{'w:ilvl':str(i)})
            for tag,attrs in [('w:start',{'w:val':'1'}),('w:numFmt',{'w:val':'decimal'}),('w:pStyle',{'w:val':f'Heading{i+1}'}),('w:lvlText',{'w:val':'.'.join('%'+str(n+1) for n in range(i+1))}),('w:suff',{'w:val':'space'})]:level.append(self.el(tag,attrs))
            abstract.append(level)
        root.append(abstract);num=self.el('w:num',{'w:numId':'70'});num.append(self.el('w:abstractNumId',{'w:val':'70'}));root.append(num)
        for i in range(3):
            np=self.el('w:numPr');np.append(self.el('w:ilvl',{'w:val':str(i)}));np.append(self.el('w:numId',{'w:val':'70'}));self.doc.styles[f'Heading {i+1}'].element.get_or_add_pPr().append(np)
    def field(self,p,instr,cached=''):
        r=p.add_run();r._r.append(self.el('w:fldChar',{'w:fldCharType':'begin'}))
        tx=self.el('w:instrText');tx.set(qn('xml:space'),'preserve');tx.text=' '+instr+' ';r._r.append(tx)
        r._r.append(self.el('w:fldChar',{'w:fldCharType':'separate'}));p.add_run(cached)
        p.add_run()._r.append(self.el('w:fldChar',{'w:fldCharType':'end'}))
    def heading(self,text,level=2,page=False):
        p=self.doc.add_paragraph(text.upper() if level==1 else text,f'Heading {level}')
        if page:p.paragraph_format.page_break_before=True
        return p
    def p(self,text='',lead=None,style=None,keep=False):
        p=self.doc.add_paragraph(style=style)
        if lead:p.add_run(lead+' ').bold=True
        p.add_run(text);p.paragraph_format.keep_with_next=keep
        return p
    def label(self,text):
        p=self.p(text);p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
        return p
    def page(self):self.doc.add_page_break()
    def field_caption(self,kind,title):
        seq=self.figures if kind=='Hình' else self.tables;seq.append(title);i=len(seq);bm=f'{"fig" if kind=="Hình" else "tbl"}{i}'
        p=self.doc.add_paragraph(style='Caption');p.add_run(kind+' ')
        self.bm+=1;p._p.append(self.el('w:bookmarkStart',{'w:id':str(self.bm),'w:name':bm}))
        self.field(p,'SEQ '+('Hinh' if kind=='Hình' else 'Bang')+' \\* ARABIC',str(i))
        p._p.append(self.el('w:bookmarkEnd',{'w:id':str(self.bm)}));p.add_run('  '+title)
        return bm
    def figure(self,name,title,intro,notes=None,width=16):
        idx=len(self.figures)+1
        p=self.p(intro+' ');p.add_run('Hình ');self.field(p,f'REF fig{idx} \\h',str(idx));p.add_run(' thể hiện các đối tượng và quan hệ trong phạm vi này.')
        p.paragraph_format.keep_with_next=True
        p=self.doc.add_paragraph();p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True;p.paragraph_format.space_after=Pt(3)
        pic=p.add_run().add_picture(str(ROOT/'figures'/f'{name}.png'),width=Cm(width));pic._inline.docPr.set('descr',title)
        self.field_caption('Hình',title)
        if notes:self.p(notes)
    def table(self,headers,rows,widths,title):
        idx=len(self.tables)+1
        p=self.p('Nội dung được tổng hợp trong Bảng ');self.field(p,f'REF tbl{idx} \\h',str(idx));p.add_run('.');p.paragraph_format.keep_with_next=True
        self.field_caption('Bảng',title)
        t=self.doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
        borders=self.el('w:tblBorders')
        for edge in ['top','left','bottom','right','insideH','insideV']:borders.append(self.el('w:'+edge,{'w:val':'single','w:sz':'4','w:color':'D9D9D9'}))
        t._tbl.tblPr.append(borders)
        for i,w in enumerate(widths):t.columns[i].width=Cm(w)
        for row_idx,values in enumerate([headers]+rows):
            row=t.rows[0] if row_idx==0 else t.add_row()
            row._tr.get_or_add_trPr().append(self.el('w:cantSplit'))
            if row_idx==0:row._tr.get_or_add_trPr().append(self.el('w:tblHeader'))
            for i,value in enumerate(values):
                cell=row.cells[i];cell.width=Cm(widths[i]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                margins=self.el('w:tcMar')
                for edge in ['top','left','bottom','right']:margins.append(self.el('w:'+edge,{'w:w':'85','w:type':'dxa'}))
                cell._tc.get_or_add_tcPr().append(margins)
                if row_idx==0:cell._tc.get_or_add_tcPr().append(self.el('w:shd',{'w:fill':'EDEDED'}))
                p=cell.paragraphs[0];p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1.05
                p.alignment=WD_ALIGN_PARAGRAPH.CENTER if row_idx==0 else WD_ALIGN_PARAGRAPH.LEFT
                run=p.add_run(value);run.font.size=Pt(11);run.bold=row_idx==0
        self.p('').paragraph_format.space_after=Pt(0)
    def spec(self,uc):
        start=len(self.doc.paragraphs)
        self.heading(uc['name'],3)
        for label,key in [('Mục tiêu','goal'),('Tác nhân','actors'),('Điều kiện trước','pre')]:
            p=self.label('');p.add_run(label+': ').bold=True;p.add_run(uc[key]);p.paragraph_format.space_after=Pt(3);p.paragraph_format.keep_with_next=True
        self.label('Luồng chính').runs[0].bold=True
        # Each use case restarts a true Word numbered list.
        root=self.doc.part.numbering_part.element;numid=100+len(root.findall(qn('w:num')))
        num=self.el('w:num',{'w:numId':str(numid)});num.append(self.el('w:abstractNumId',{'w:val':'7'}));override=self.el('w:lvlOverride',{'w:ilvl':'0'});override.append(self.el('w:startOverride',{'w:val':'1'}));num.append(override);root.append(num)
        for step in uc['flow']:
            p=self.doc.add_paragraph(step,'List Number');p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1.05
            np=self.el('w:numPr');np.append(self.el('w:ilvl',{'w:val':'0'}));np.append(self.el('w:numId',{'w:val':str(numid)}));p._p.get_or_add_pPr().append(np)
        for label,key in [('Ngoại lệ và nhánh khác','alt'),('Điều kiện sau','post')]:
            p=self.label('');p.add_run(label+': ').bold=True;p.add_run(uc[key]);p.paragraph_format.space_after=Pt(4)
        for p in self.doc.paragraphs[start:-1]:p.paragraph_format.keep_with_next=True
    def save(self,path):self.doc.save(path)
